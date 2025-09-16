from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pandas as pd

def remove_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top','left','bottom','right','insideH','insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'nil')
        tcBorders.append(border)
    tcPr.append(tcBorders)

def set_cell_width(cell, width):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(width))
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)

doc = Document()
df = pd.read_csv("../Desktop/r17/summary_metrics.csv")

cols = ["Station", "Timescale", "Std Ref", "Model","Std Model","RMSE","Corr","CRMSE","MAE","MAPE"]
table = doc.add_table(rows=1, cols=len(cols))
table.style = None

cell_widths = [2000, 2000, 1500, 2000, 2000, 1200, 1200, 1200, 1200, 1200]

for i, col in enumerate(cols):
    table.rows[0].cells[i].text = col
    remove_cell_borders(table.rows[0].cells[i])
    set_cell_width(table.rows[0].cells[i], cell_widths[i])

last_station = None
station_start_idx = None

for station, station_group in df.groupby("station"):
    station_start_idx = len(table.rows)
    last_spi = None
    timescale_start_idx = None

    for spi, spi_group in station_group.groupby("spi"):
        timescale_start_idx = len(table.rows)

        for _, r in spi_group.iterrows():
            row = table.add_row().cells
            row[0].text = str(station)      # temporary, will merge later
            row[1].text = str(spi)          # temporary, will merge later
            row[2].text = str(r["std_ref"])
            for i, col in enumerate(["model","std_model","rmse","corr","crmse","mae","mape"], start=3):
                row[i].text = str(r[col])
            for idx, cell in enumerate(row):
                remove_cell_borders(cell)
                set_cell_width(cell, cell_widths[idx])

        # merge Timescale column for this spi group
        end_idx = len(table.rows)-1
        if end_idx > timescale_start_idx:
            table.cell(timescale_start_idx,1).merge(table.cell(end_idx,1))

    # merge Station column for this station group
    end_station_idx = len(table.rows)-1
    if end_station_idx > station_start_idx:
        table.cell(station_start_idx,0).merge(table.cell(end_station_idx,0))

    # add horizontal line separator after each station
    sep_row = table.add_row().cells
    sep_row[0].merge(sep_row[-1])
    sep_row[0].text = ""
    tc = sep_row[0]._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), '000000')
    tcBorders.append(bottom)
    tcPr.append(tcBorders)

doc.save("Results/SPI_results.docx")
print("Created SPI_results.docx")
